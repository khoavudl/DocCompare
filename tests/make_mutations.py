import shutil
import sys
from pathlib import Path
import docx

sys.path.insert(0, str(Path(__file__).parent.parent))
from core.normalizer import docx_to_pdf

SAMPLE_01_DOCX = Path(__file__).parent / "fixtures" / "sample_01" / "approved.docx"
MUTATIONS_DIR = Path(__file__).parent / "fixtures" / "mutations"


def apply_mutation(mut_id: str, doc: docx.Document) -> None:
    """Applies mutation logic to a copy of the Document."""
    if mut_id == "M1":
        # M1: Đổi 1 chữ số: "8%" → "9%" (Điều 4.6)
        for p in doc.paragraphs:
            if "8%" in p.text:
                p.text = p.text.replace("8%", "9%")
                break

    elif mut_id == "M2":
        # M2: Đổi số tiền: "500.000VNĐ" → "800.000VNĐ" (Điều 4.2)
        for p in doc.paragraphs:
            if "500.000VNĐ" in p.text:
                p.text = p.text.replace("500.000VNĐ", "800.000VNĐ")
                break

    elif mut_id == "M3":
        # M3: Xoá cả câu "Bên B phải chịu các chi phí phát sinh liên quan đến việc chuyển khoản lại."
        for p in doc.paragraphs:
            target = "Bên B phải chịu các chi phí phát sinh liên quan đến việc chuyển khoản lại."
            if target in p.text:
                p.text = p.text.replace(target, "").strip()
                break

    elif mut_id == "M4":
        # M4: Thêm câu mới vào cuối Điều 7.2
        doc.paragraphs[216].text += " This is an additional confidentiality clause."

    elif mut_id == "M5":
        # M5: Đổi "Youtube" → "Tiktok" (Điều 2.1)
        doc.paragraphs[79].text = doc.paragraphs[79].text.replace("Youtube", "Tiktok")

    elif mut_id == "M6":
        # M6: Đổi 1 chữ ở dòng cuối cùng của một trang ("budgeting." ở cuối trang 2)
        doc.paragraphs[83].text = doc.paragraphs[83].text.replace("budgeting.", "finance.")

    elif mut_id == "M7":
        # M7: Đổi 1 chữ ở dòng đầu tiên của một trang ("ADIGIT" ở đầu trang 2)
        doc.paragraphs[45].text = doc.paragraphs[45].text.replace("ADIGIT", "BDIGIT")

    elif mut_id == "M8":
        # M8: Đổi "SONI" → "Soni"
        doc.paragraphs[37].text = doc.paragraphs[37].text.replace("SONI", "Soni")

    elif mut_id == "M9":
        # M9: Đổi số tài khoản "612345688" → "612345699"
        doc.paragraphs[52].text = doc.paragraphs[52].text.replace("612345688", "612345699")

    elif mut_id == "M10":
        # M10: Đổi dấu: "Bình Hòa" → "Bình Hoà"
        doc.paragraphs[35].text = doc.paragraphs[35].text.replace("Bình Hòa", "Bình Hoà")

    elif mut_id == "M11":
        # M11: Đổi một số đứng riêng thành một dòng (số "3" ở trang 2)
        doc.paragraphs[75].insert_paragraph_before("3")

    elif mut_id == "M12":
        # M12: Không đổi gì (bản đối chứng)
        pass


def generate_all_mutations():
    MUTATIONS_DIR.mkdir(parents=True, exist_ok=True)
    mut_ids = [f"M{i}" for i in range(1, 13)]

    for mut_id in mut_ids:
        target_dir = MUTATIONS_DIR / mut_id
        target_dir.mkdir(parents=True, exist_ok=True)
        docx_out = target_dir / "mutated.docx"

        doc = docx.Document(str(SAMPLE_01_DOCX))
        apply_mutation(mut_id, doc)
        doc.save(str(docx_out))

        print(f"[{mut_id}] Converting to PDF...")
        pdf_out = docx_to_pdf(str(docx_out), str(target_dir))
        # Chuẩn hoá tên file pdf thành fake_wesign.pdf
        dest_pdf = target_dir / "fake_wesign.pdf"
        if Path(pdf_out).resolve() != dest_pdf.resolve():
            shutil.move(pdf_out, dest_pdf)

        print(f"[{mut_id}] Done -> {dest_pdf}")


if __name__ == "__main__":
    generate_all_mutations()
